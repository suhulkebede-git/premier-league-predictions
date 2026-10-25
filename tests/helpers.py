"""Builders for synthetic prediction sheets used across the test suite."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from plpredict.models import AWARD_CATEGORIES
from plpredict.names import Aliases

CLUBS = [
    "AFC Bournemouth", "Arsenal", "Aston Villa", "Brentford", "Brighton & Hove Albion",
    "Chelsea", "Coventry City", "Crystal Palace", "Everton", "Fulham", "Hull City",
    "Ipswich Town", "Leeds United", "Liverpool", "Manchester City", "Manchester United",
    "Newcastle United", "Nottingham Forest", "Sunderland", "Tottenham Hotspur",
]
# Last season: 17 surviving clubs + 3 relegated (not in this season).
LAST_SEASON = [
    ("Arsenal C", 85), ("Manchester City", 78), ("Manchester United", 71),
    ("Aston Villa", 65), ("Liverpool", 60), ("Bournemouth", 57), ("Sunderland", 54),
    ("Brighton", 53), ("Brentford", 53), ("Chelsea", 52), ("Fulham", 52),
    ("Newcastle United", 49), ("Everton", 49), ("Leeds United", 47),
    ("Crystal Palace", 45), ("Nottingham Forest", 44), ("Tottenham Hotspur", 41),
    ("West Ham United", 39), ("Burnley", 22), ("Wolverhampton Wanderers", 20),
]
DEFAULT_PICKS = [
    "Bruno F", "Jakirovich", "6 points", "Spurs", "Sunderland", "Leeds", "Brunooooo",
    "Raya", "Crystal Palace or Brighton", "Coventry", "Big Erl", "Arsenal", "Hull",
    "Hull City", "Hull", "Man City", "Arsenal", "Hell nah", "Yes", "No",
]


def test_aliases() -> Aliases:
    return Aliases(
        clubs={
            "arsenal c": "Arsenal", "bournemouth": "AFC Bournemouth",
            "brighton": "Brighton & Hove Albion", "spurs": "Tottenham Hotspur",
            "leeds": "Leeds United", "coventry": "Coventry City", "hull": "Hull City",
            "man city": "Manchester City", "raya": "Arsenal", "man utd": "Manchester United",
            "nott'm forest": "Nottingham Forest",
        },
        players={"bruno f": "Bruno Fernandes", "brunooooo": "Bruno Fernandes",
                 "big erl": "Erling Haaland"},
        managers={"jakirovich": "Sergej Jakirovic"},
        answers={"yes": True, "no": False, "hell nah": False},
    )


def write_sheet(
    path: Path,
    people: list[tuple[str, list[str] | None, list[object] | None]],
    *,
    season_clubs: list[str] | None = None,
) -> Path:
    """Write a workbook in the group's layout.

    ``people`` is a list of (name, club order or None for blank, picks or None
    for DEFAULT_PICKS); blocks fill left then right, 24 rows apart, from row 26.
    """
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "LAST TIME OUT", "THE NEW SEASON"
    ws["A2"], ws["B2"], ws["C2"] = "Club", "Points", "Clubs"
    for i, club in enumerate(season_clubs or CLUBS):
        ws.cell(row=3 + i, column=3, value=club)
    for i, (club, pts) in enumerate(LAST_SEASON):
        ws.cell(row=3 + i, column=1, value=club)
        ws.cell(row=3 + i, column=2, value=pts)
    for n, (name, order, picks) in enumerate(people):
        top, col0 = 26 + 24 * (n // 2), 1 if n % 2 == 0 else 6
        ws.cell(row=top, column=col0, value="Name")
        ws.cell(row=top, column=col0 + 1, value=name)
        for j, header in enumerate(("PREDIC", "CLUB", "INDIVIDUAL AWARDS")):
            ws.cell(row=top + 1, column=col0 + j, value=header)
        for i, category in enumerate(AWARD_CATEGORIES):
            r = top + 2 + i
            ws.cell(row=r, column=col0, value=float(i + 1) if i < 3 else i + 1)
            ws.cell(row=r, column=col0 + 2, value=category.label)
            if order is not None:
                ws.cell(row=r, column=col0 + 1, value=order[i])
                ws.cell(row=r, column=col0 + 3, value=(picks or DEFAULT_PICKS)[i])
    wb.save(path)
    return path
